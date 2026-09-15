#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Tái tạo trang bìa BaoCao_SD16_SAOClub.docx theo form của project Graceful SD-56.

Cover paragraph gốc có 3 <w:r>:
  - run[0]: 1 rect (text box 'Vũ Quang Huy' ở góc)
  - run[1]: 1 rect (text box 'PH61463' ở góc)
  - run[2]: 1 <w:pict> với <v:group> chứa:
      - 129 <v:rect> (text box rời rạc nội dung cover)
      - 1 <v:shape> (Picture - logo FPT)

Cách làm:
  - Xóa run[0] và run[1]
  - Trong run[2]: xóa tất cả <v:rect>, giữ <v:group> và <v:shape> (logo)
  - Chèn các paragraph mới sau cover paragraph
"""

import sys
import io
from copy import deepcopy
from docx import Document
from lxml import etree

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
XML_NS = "http://www.w3.org/XML/1998/namespace"


def qn_w(tag):
    return f"{{{W_NS}}}{tag}"


def qn_v(tag):
    return f"{{{V_NS}}}{tag}"


def make_run(text, *, bold=False, size=26):
    r = etree.Element(qn_w("r"))
    rPr = etree.SubElement(r, qn_w("rPr"))
    rFonts = etree.SubElement(rPr, qn_w("rFonts"))
    rFonts.set(qn_w("ascii"), "Times New Roman")
    rFonts.set(qn_w("hAnsi"), "Times New Roman")
    rFonts.set(qn_w("cs"), "Times New Roman")
    if bold:
        etree.SubElement(rPr, qn_w("b"))
        etree.SubElement(rPr, qn_w("bCs"))
    sz = etree.SubElement(rPr, qn_w("sz"))
    sz.set(qn_w("val"), str(size))
    szCs = etree.SubElement(rPr, qn_w("szCs"))
    szCs.set(qn_w("val"), str(size))
    t = etree.SubElement(r, qn_w("t"))
    t.set(f"{{{XML_NS}}}space", "preserve")
    t.text = text
    return r


def make_label_value(label, value, *, size=26):
    p = etree.Element(qn_w("p"))
    pPr = etree.SubElement(p, qn_w("pPr"))
    spacing = etree.SubElement(pPr, qn_w("spacing"))
    spacing.set(qn_w("after"), "120")
    spacing.set(qn_w("line"), "360")
    spacing.set(qn_w("lineRule"), "auto")
    tabs = etree.SubElement(pPr, qn_w("tabs"))
    tab = etree.SubElement(tabs, qn_w("tab"))
    tab.set(qn_w("val"), "left")
    tab.set(qn_w("pos"), "3402")
    ind = etree.SubElement(pPr, qn_w("ind"))
    ind.set(qn_w("left"), "2268")
    p.append(make_run(label, bold=True, size=size))
    p.append(make_run(" : ", bold=True, size=size))
    p.append(make_run(value, bold=False, size=size))
    return p


def make_sv_paragraph(name, code, *, size=26):
    p = etree.Element(qn_w("p"))
    pPr = etree.SubElement(p, qn_w("pPr"))
    spacing = etree.SubElement(pPr, qn_w("spacing"))
    spacing.set(qn_w("after"), "60")
    spacing.set(qn_w("line"), "360")
    spacing.set(qn_w("lineRule"), "auto")
    tabs = etree.SubElement(pPr, qn_w("tabs"))
    tab = etree.SubElement(tabs, qn_w("tab"))
    tab.set(qn_w("val"), "left")
    tab.set(qn_w("pos"), "5103")
    ind = etree.SubElement(pPr, qn_w("ind"))
    ind.set(qn_w("left"), "2268")
    p.append(make_run(name, bold=False, size=size))
    p.append(make_run("\t", bold=False, size=size))
    p.append(make_run(code, bold=False, size=size))
    return p


def make_centered_paragraph(text, *, bold=False, size=26, spacing_after=240, line=360):
    p = etree.Element(qn_w("p"))
    pPr = etree.SubElement(p, qn_w("pPr"))
    spacing = etree.SubElement(pPr, qn_w("spacing"))
    spacing.set(qn_w("after"), str(spacing_after))
    spacing.set(qn_w("line"), str(line))
    spacing.set(qn_w("lineRule"), "auto")
    jc = etree.SubElement(pPr, qn_w("jc"))
    jc.set(qn_w("val"), "center")
    if text:
        p.append(make_run(text, bold=bold, size=size))
    return p


def make_blank_paragraph(spacing_after=240, line=360):
    p = etree.Element(qn_w("p"))
    pPr = etree.SubElement(p, qn_w("pPr"))
    spacing = etree.SubElement(pPr, qn_w("spacing"))
    spacing.set(qn_w("after"), str(spacing_after))
    spacing.set(qn_w("line"), str(line))
    spacing.set(qn_w("lineRule"), "auto")
    return p


def main():
    doc_path = "BaoCao_SD16_SAOClub.docx"
    doc = Document(doc_path)
    body = doc.element.body

    first_p = list(body)[0]
    assert first_p.tag == qn_w("p"), f"first body child is not <w:p>: {first_p.tag}"

    # 1. Xóa run[0] và run[1] (text box tên SV ở góc)
    runs = first_p.findall(qn_w("r"))
    print(f"Total <w:r> in cover: {len(runs)}")

    textbox_runs_to_remove = []
    logo_run = None
    for r in runs:
        pict = r.find(qn_w("pict"))
        if pict is None:
            continue
        imagedata = pict.find(".//" + qn_v("imagedata"))
        if imagedata is not None:
            logo_run = r
        else:
            textbox_runs_to_remove.append(r)

    print(f"Text-box runs to remove: {len(textbox_runs_to_remove)}")
    print(f"Logo run found: {logo_run is not None}")

    for r in textbox_runs_to_remove:
        first_p.remove(r)

    # 2. Trong logo_run: xóa tất cả <v:rect> con, giữ <v:shape> (ảnh) và <v:group>
    if logo_run is not None:
        pict = logo_run.find(qn_w("pict"))
        rects = pict.findall(".//" + qn_v("rect"))
        print(f"  Removing {len(rects)} <v:rect> from logo_run's pict")
        for rect in rects:
            parent = rect.getparent()
            if parent is not None:
                parent.remove(rect)

    # 3. Tạo các paragraph mới cho cover
    new_paragraphs = []

    # Logo FPT: copy logo_run sang paragraph riêng, căn giữa
    if logo_run is not None:
        logo_p = etree.Element(qn_w("p"))
        logo_pPr = etree.SubElement(logo_p, qn_w("pPr"))
        logo_spacing = etree.SubElement(logo_pPr, qn_w("spacing"))
        logo_spacing.set(qn_w("after"), "240")
        logo_spacing.set(qn_w("line"), "240")
        logo_spacing.set(qn_w("lineRule"), "auto")
        logo_jc = etree.SubElement(logo_pPr, qn_w("jc"))
        logo_jc.set(qn_w("val"), "center")
        logo_p.append(deepcopy(logo_run))
        new_paragraphs.append(logo_p)

        # Xóa logo_run cũ trong first_p (đã copy sang paragraph mới)
        first_p.remove(logo_run)

    # Reset pPr của cover paragraph cũ
    pPr = first_p.find(qn_w("pPr"))
    if pPr is not None:
        for child in list(pPr):
            pPr.remove(child)
        sp = etree.SubElement(pPr, qn_w("spacing"))
        sp.set(qn_w("after"), "0")
        sp.set(qn_w("line"), "240")
        sp.set(qn_w("lineRule"), "auto")

    # TRƯỜNG CAO ĐẲNG FPT POLYTECHNIC
    new_paragraphs.append(
        make_centered_paragraph("TRƯỜNG CAO ĐẲNG FPT POLYTECHNIC", bold=True, size=28, spacing_after=120)
    )

    # BÁO CÁO DỰ ÁN TỐT NGHIỆP
    new_paragraphs.append(
        make_centered_paragraph("BÁO CÁO DỰ ÁN TỐT NGHIỆP", bold=True, size=32, spacing_after=120)
    )

    # Tên đề tài
    new_paragraphs.append(
        make_centered_paragraph("Xây dựng website bán máy tính SAOClub", bold=True, size=34, spacing_after=480)
    )

    # Khoảng trống
    new_paragraphs.append(make_blank_paragraph(spacing_after=240))

    # Thông tin chung
    new_paragraphs.append(make_label_value("Giảng viên hướng dẫn", "Nguyễn Quang Hà"))
    new_paragraphs.append(make_label_value("Chuyên ngành", "Phát triển phần mềm"))
    new_paragraphs.append(make_label_value("Nhóm thực hiện", "SD-16"))

    # Sinh viên thực hiện label
    sv_header = etree.Element(qn_w("p"))
    sv_header_pPr = etree.SubElement(sv_header, qn_w("pPr"))
    sv_header_spacing = etree.SubElement(sv_header_pPr, qn_w("spacing"))
    sv_header_spacing.set(qn_w("after"), "120")
    sv_header_spacing.set(qn_w("line"), "360")
    sv_header_spacing.set(qn_w("lineRule"), "auto")
    sv_header_ind = etree.SubElement(sv_header_pPr, qn_w("ind"))
    sv_header_ind.set(qn_w("left"), "2268")
    sv_header.append(make_run("Sinh viên thực hiện", bold=True, size=26))
    sv_header.append(make_run(" :", bold=True, size=26))
    new_paragraphs.append(sv_header)

    students = [
        ("Lê Huy Đỗ", "PH60048"),
        ("Lê Anh Ngữ", "PH60049"),
        ("Nghiêm Việt Anh", "PH59725"),
        ("Nguyễn Thành Đạt", "PH60102"),
        ("Nguyễn Xuân Việt", "PH59601"),
        ("Vũ Quang Huy", "PH61463"),
    ]
    for name, code in students:
        new_paragraphs.append(make_sv_paragraph(name, code))

    # Khoảng trống + Hà Nội - 2026 + page break
    new_paragraphs.append(make_blank_paragraph(spacing_after=480))

    ha_noi_p = make_centered_paragraph("Hà Nội - 2026", bold=True, size=26, spacing_after=0)
    page_break_run = etree.SubElement(ha_noi_p, qn_w("r"))
    rPr = etree.SubElement(page_break_run, qn_w("rPr"))
    rFonts = etree.SubElement(rPr, qn_w("rFonts"))
    rFonts.set(qn_w("ascii"), "Times New Roman")
    rFonts.set(qn_w("hAnsi"), "Times New Roman")
    br = etree.SubElement(page_break_run, qn_w("br"))
    br.set(qn_w("type"), "page")
    new_paragraphs.append(ha_noi_p)

    # 4. Chèn các paragraph mới ngay sau first_p
    first_p_index = list(body).index(first_p)
    for i, p in enumerate(new_paragraphs, start=1):
        body.insert(first_p_index + i, p)

    # 5. Lưu file
    doc.save(doc_path)
    print(f"\nSaved: {doc_path}")
    print(f"Inserted {len(new_paragraphs)} new paragraphs after cover.")

    # 6. Verify
    doc2 = Document(doc_path)
    print(f"\n=== Verification ===")
    print(f"Total paragraphs: {len(doc2.paragraphs)}")

    body2 = doc2.element.body
    children = list(body2)

    # First child = empty cover (now empty)
    # Subsequent children = new paragraphs + SDT + MỤC LỤC etc.
    print("\n=== First 25 body children ===")
    for i, child in enumerate(children[:25]):
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            texts = child.findall('.//' + qn_w('t'))
            text = ''.join(t.text or '' for t in texts).strip()
            rects = child.findall('.//' + qn_v('rect'))
            images = child.findall('.//' + qn_v('imagedata'))
            print(f"  [{i}] <p> text='{text[:50]}' rects={len(rects)} images={len(images)}")
        else:
            print(f"  [{i}] <{tag}>")

    # Verify total rects in cover area (before MỤC LỤC SDT)
    cover_rects = 0
    cover_images = 0
    for child in children:
        if child.tag == qn_w('p'):
            rects = child.findall('.//' + qn_v('rect'))
            images = child.findall('.//' + qn_v('imagedata'))
            # Stop at SDT-containing paragraph
            if child.find('.//' + qn_w('sdt')) is not None:
                break
            cover_rects += len(rects)
            cover_images += len(images)
    print(f"\nCover area: {cover_rects} rects, {cover_images} images (should be 0 rects, 1 image)")

    # Verify page break
    for i, p in enumerate(doc2.paragraphs):
        if 'Hà Nội' in p.text:
            pbs = [b for b in p._p.findall(qn_w('r') + '/' + qn_w('br')) if b.get(qn_w('type')) == 'page']
            print(f"Page break after 'Hà Nội - 2026' (paragraph {i}): {len(pbs)}")
            break

    print("\n=== Cover content ===")
    for i, p in enumerate(doc2.paragraphs[:22]):
        txt = p.text[:70]
        print(f"  [{i}] '{txt}'")


if __name__ == "__main__":
    main()
