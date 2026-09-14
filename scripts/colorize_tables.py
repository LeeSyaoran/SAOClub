"""Tô màu tất cả các bảng trong BaoCao_SD16_SAOClub.docx"""
import sys
from docx import Document
from docx.oxml.ns import qn, nsmap
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy

sys.stdout.reconfigure(encoding='utf-8')

INPUT = 'BaoCao_SD16_SAOClub.docx'
OUTPUT = 'BaoCao_SD16_SAOClub.docx'

# Bảng màu (phối màu theo theme tài liệu)
HEADER_BG = '1F3864'           # Xanh navy đậm - header
HEADER_TEXT = 'FFFFFF'         # Trắng
ALT_ROW_BG = 'F2F2F2'          # Xám nhạt - dòng xen kẽ
NORMAL_BG = 'FFFFFF'           # Trắng
BORDER_COLOR = '000000'        # Đen

doc = Document(INPUT)
total_tables = len(doc.tables)

# Style cho header & body
def set_cell_background(cell, hex_color):
    """Set background color for a cell"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = tc_pr.makeelement(qn('w:shd'), {})
        tc_pr.append(shd)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)

def set_cell_text_color(cell, hex_color, bold=False):
    """Set text color and bold for all paragraphs in cell"""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(hex_color)
            if bold:
                run.font.bold = True

def set_cell_alignment(cell, alignment='center'):
    """Set cell content alignment"""
    align_map = {
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
    }
    for paragraph in cell.paragraphs:
        if alignment in align_map:
            paragraph.alignment = align_map[alignment]

def set_table_borders(table):
    """Set borders for entire table"""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = tbl.makeelement(qn('w:tblPr'), {})
        tbl.insert(0, tblPr)

    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is None:
        tblBorders = tblPr.makeelement(qn('w:tblBorders'), {})
        tblPr.append(tblBorders)

    # Xóa borders cũ
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        existing = tblBorders.find(qn(f'w:{border_name}'))
        if existing is not None:
            tblBorders.remove(existing)

    # Thêm borders mới
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = tblBorders.makeelement(qn(f'w:{border_name}'), {})
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '6')  # 0.75pt
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), BORDER_COLOR)
        tblBorders.append(border)

# Tô màu từng bảng
processed = 0
skipped = 0
for idx, table in enumerate(doc.tables):
    try:
        rows = table.rows
        if len(rows) == 0:
            skipped += 1
            continue

        # Kiểm tra nếu là bảng đặc biệt (1 row 1 col) - bỏ qua
        if len(rows) == 1 and len(rows[0].cells) == 1:
            skipped += 1
            continue
        # Bảng 1x2 cũng có thể là layout box - check content
        if len(rows) == 1 and len(rows[0].cells) == 2:
            cell_texts = [c.text.strip() for c in rows[0].cells]
            if all(len(t) < 5 for t in cell_texts):
                skipped += 1
                continue

        set_table_borders(table)

        # Header row (row 0) - xanh navy đậm, chữ trắng, bold
        for cell in rows[0].cells:
            set_cell_background(cell, HEADER_BG)
            set_cell_text_color(cell, HEADER_TEXT, bold=True)
            set_cell_alignment(cell, 'center')

        # Body rows - alternate colors
        for i, row in enumerate(rows[1:], start=1):
            bg_color = ALT_ROW_BG if i % 2 == 1 else NORMAL_BG
            for cell in row.cells:
                set_cell_background(cell, bg_color)
                # Body text mặc định màu đen
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor.from_string('000000')

        processed += 1
    except Exception as e:
        print(f'Lỗi bảng {idx}: {e}')
        skipped += 1

print(f'Đã xử lý: {processed}/{total_tables} bảng (bỏ qua: {skipped})')

doc.save(OUTPUT)
print(f'Đã lưu: {OUTPUT}')
